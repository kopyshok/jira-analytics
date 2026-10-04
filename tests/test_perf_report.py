"""Чтение замеров за период: узкие места, ряд для графика, медленные и вывод."""
from datetime import datetime, timedelta

import pytest

from app.core.perf import HIST_SIZE, hist_index, p95_from_hist
from app.models.perf import PerfMinute, PerfServerSnapshot, PerfSlowRequest
from app.services.perf_report import overview, section_label, verdict

NOW = datetime(2026, 10, 4, 12, 30)
STUB_USER_ID = "00000000-0000-0000-0000-000000000000"


def _minute(db, at, route, durations, *, method="GET", db_count=0, db_ms=0.0, errors=0):
    hist = [0] * HIST_SIZE
    for d in durations:
        hist[hist_index(d)] += 1
    db.add(PerfMinute(
        minute=at, method=method, route=route, count=len(durations),
        total_ms=float(sum(durations)), max_ms=float(max(durations)), errors_5xx=errors,
        db_count=db_count, db_ms=db_ms, **{f"h{i}": n for i, n in enumerate(hist)},
    ))


def _snap(db, at, host, proc, cpu_count=4):
    db.add(PerfServerSnapshot(
        at=at, host_cpu_percent=host, host_memory_percent=50, process_cpu_percent=proc,
        process_memory_mb=300, cpu_count=cpu_count, threads=20, db_pool_in_use=2,
        db_pool_size=40, requests_in_flight=3,
    ))


def _slow(db, at, *, duration=3000.0, db_ms=0.0, cpu_ms=0.0, route="/api/v1/backlog/{item_id}",
          user_id=None):
    db.add(PerfSlowRequest(
        at=at, method="GET", route=route, path="/api/v1/backlog/1", query="team=A",
        status_code=200, duration_ms=duration, db_count=4, db_ms=db_ms, cpu_ms=cpu_ms,
        user_id=user_id, top_queries=[{"ms": db_ms, "sql": "SELECT * FROM backlog_items"}],
    ))


@pytest.mark.parametrize("kw, expected", [
    (dict(host_cpu=95, process_share=10, process_core=40, duration_ms=3000, db_ms=0, cpu_ms=0),
     "other_load"),
    (dict(host_cpu=40, process_share=10, process_core=40, duration_ms=3000, db_ms=2000, cpu_ms=0),
     "database"),
    (dict(host_cpu=40, process_share=10, process_core=40, duration_ms=3000, db_ms=0, cpu_ms=2700),
     "our_code"),
    (dict(host_cpu=40, process_share=25, process_core=95, duration_ms=3000, db_ms=0, cpu_ms=0),
     "our_code"),
    (dict(host_cpu=40, process_share=5, process_core=20, duration_ms=3000, db_ms=100, cpu_ms=100),
     "waiting"),
    (dict(host_cpu=None, process_share=None, process_core=None, duration_ms=3000, db_ms=0,
          cpu_ms=0), "waiting"),
    # Ожидание подключения из пула во время в базе не входит — отдельное правило.
    (dict(host_cpu=40, process_share=5, process_core=20, duration_ms=3000, db_ms=100, cpu_ms=100,
          pool_in_use=40, pool_size=40), "db_pool"),
    (dict(host_cpu=40, process_share=5, process_core=20, duration_ms=3000, db_ms=100, cpu_ms=100,
          pool_in_use=39, pool_size=40), "waiting"),
])
def test_verdict(kw, expected):
    assert verdict(**kw) == expected


def test_section_label_by_route():
    assert section_label("/api/v1/backlog/{item_id}") == "Целевые задачи"
    assert section_label("/api/v1/resource-planning/plans") == "Ресурсное планирование"
    assert section_label("/api/v1/whatever") == "whatever"


def test_bottlenecks_sum_rows_of_all_flushes_and_sort_by_total_time(db_session):
    m = NOW - timedelta(minutes=10)
    # Один шаблон пути, два сброса в одну минуту — строки суммируются.
    _minute(db_session, m, "/api/v1/backlog/{item_id}", [100] * 10, db_count=20, db_ms=500)
    _minute(db_session, m, "/api/v1/backlog/{item_id}", [100] * 9 + [900], db_count=20, db_ms=500,
            errors=1)
    _minute(db_session, m, "/api/v1/planning/scenarios", [5000])
    # Вне периода «час» — не считается.
    _minute(db_session, NOW - timedelta(hours=2), "/api/v1/planning/scenarios", [9000])
    db_session.commit()

    data = overview(db_session, "1h", now=NOW, slow_ms=2000, flush_seconds=60)

    first, second = data["bottlenecks"]
    assert first["route"] == "/api/v1/planning/scenarios"
    assert first["calls"] == 1 and first["total_ms"] == pytest.approx(5000)
    assert second["route"] == "/api/v1/backlog/{item_id}"
    assert second["section"] == "Целевые задачи"
    assert second["calls"] == 20
    assert second["avg_ms"] == pytest.approx(140)
    assert second["max_ms"] == pytest.approx(900)
    assert second["errors_5xx"] == 1
    assert second["db_avg_count"] == pytest.approx(2)
    assert second["db_share"] == pytest.approx(1000 / 2800)
    hist = [0] * HIST_SIZE
    hist[hist_index(100)] = 19
    hist[hist_index(900)] = 1
    assert second["p95_ms"] == pytest.approx(p95_from_hist(hist, 900))
    assert data["totals"]["requests"] == 21
    assert data["totals"]["errors_5xx"] == 1


def test_series_has_bucket_per_minute_with_load(db_session):
    m = NOW - timedelta(minutes=5)
    _minute(db_session, m, "/api/v1/backlog/{item_id}", [100, 200])
    _snap(db_session, m + timedelta(seconds=30), host=60, proc=80, cpu_count=4)
    db_session.commit()

    data = overview(db_session, "1h", now=NOW, slow_ms=2000, flush_seconds=60)

    assert data["bucket_minutes"] == 1
    assert len(data["series"]) == 61
    point = next(p for p in data["series"] if p["requests"])
    assert point["t"].startswith("2026-10-04T12:25:00")
    assert point["requests"] == 2
    assert point["host_cpu"] == pytest.approx(60)
    assert point["process_cpu"] == pytest.approx(20)  # 80% ядра из 4 ядер
    empty = data["series"][0]
    assert empty["requests"] == 0 and empty["p95_ms"] is None and empty["host_cpu"] is None


def test_slow_requests_get_minute_load_and_verdict(db_session):
    at = NOW - timedelta(minutes=3, seconds=20)
    _snap(db_session, NOW - timedelta(minutes=3), host=95, proc=20, cpu_count=2)  # доля 10%
    _slow(db_session, at, user_id=STUB_USER_ID)
    _slow(db_session, NOW - timedelta(minutes=50), db_ms=2500)  # снимка рядом нет
    db_session.commit()

    data = overview(db_session, "1h", now=NOW, slow_ms=2000, flush_seconds=60)

    first, second = data["slow"]
    assert first["user"] == "Test User"
    assert first["section"] == "Целевые задачи"
    assert first["host_cpu"] == pytest.approx(95)
    assert first["process_cpu"] == pytest.approx(10)
    assert first["verdict"] == "other_load"
    assert first["verdict_label"]
    assert first["verdict_reason"] == "сервер загружен на 95%, наш сервис — 10%"
    assert first["top_queries"][0]["sql"].startswith("SELECT")
    assert second["host_cpu"] is None
    assert second["verdict"] == "database"
    assert second["verdict_reason"] == "ожидание базы — 83% времени запроса"
    assert data["verdicts"] == {"other_load": 1, "database": 1}
    assert data["totals"]["slow"] == 2
    bucket = next(p for p in data["series"] if p["slow"])
    assert bucket["verdict"] in {"other_load", "database"}


def test_period_validation():
    with pytest.raises(ValueError):
        overview(None, "2h", now=NOW, slow_ms=2000, flush_seconds=60)  # type: ignore[arg-type]


def _seed_for_export(db):
    m = NOW - timedelta(minutes=5)
    _minute(db, m, "/api/v1/backlog/{item_id}", [100, 3000], db_count=4, db_ms=2600)
    _snap(db, NOW - timedelta(minutes=4), host=40, proc=30, cpu_count=4)
    _slow(db, m + timedelta(seconds=10), db_ms=2600, user_id=STUB_USER_ID)
    db.commit()


def test_markdown_report_for_developers(db_session):
    from app.services.perf_report import render_markdown

    _seed_for_export(db_session)
    data = overview(db_session, "1h", now=NOW, slow_ms=2000, flush_seconds=60)
    text = render_markdown(data, tz_offset_min=180)

    assert text.startswith("# Быстродействие сервиса")
    for part in ("## Итог", "## Узкие места", "## Медленные запросы", "## Нагрузка сервера"):
        assert part in text
    assert "`GET /api/v1/backlog/{item_id}`" in text
    assert "Целевые задачи" in text
    assert "Test User" in text
    assert "Долгая работа с базой" in text
    assert "SELECT * FROM backlog_items" in text
    assert "15:25" in text  # время сдвинуто в пояс UTC+3
    assert "UTC+03:00" in text


def test_xlsx_has_three_sheets(db_session):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.perf_report import render_xlsx, snapshots_for

    _seed_for_export(db_session)
    data = overview(db_session, "1h", now=NOW, slow_ms=2000, flush_seconds=60)
    blob = render_xlsx(data, snapshots_for(db_session, "1h", now=NOW), tz_offset_min=180)

    wb = load_workbook(BytesIO(blob))
    assert wb.sheetnames == ["Узкие места", "Медленные запросы", "Снимки сервера"]
    assert wb["Узкие места"].max_row == 2
    assert wb["Узкие места"]["A1"].value == "Раздел"
    assert wb["Медленные запросы"].max_row == 2
    assert wb["Снимки сервера"].max_row == 2
    assert wb["Медленные запросы"]["A2"].value.hour == 15


def test_verdict_reason_names_the_deciding_number():
    from types import SimpleNamespace

    from app.services.perf_report import verdict_reason

    slow = SimpleNamespace(duration_ms=3000.0, db_ms=0.0, cpu_ms=2700.0)
    snap = SimpleNamespace(host_cpu_percent=40.0, process_cpu_percent=95.0, cpu_count=4)
    assert verdict_reason("our_code", slow, snap) == (
        "наш сервис за время запроса занимал процессор на 90% одного ядра"
    )
    idle = SimpleNamespace(duration_ms=3000.0, db_ms=0.0, cpu_ms=0.0)
    assert verdict_reason("our_code", idle, snap) == (
        "наш сервис в ту минуту занимал процессор на 95% одного ядра"
    )
    assert verdict_reason("waiting", idle, None) == "процессор и база почти не были заняты"
    full = SimpleNamespace(db_pool_in_use=40, db_pool_size=40)
    assert verdict_reason("db_pool", idle, full) == "заняты все подключения к базе: 40 из 40"


def test_slow_total_is_counted_beyond_cause_sample(db_session, monkeypatch):
    """Число медленных — счётом в базе; причины — по последним SLOW_STATS_LIMIT."""
    from app.services import perf_report
    from app.services.perf_report import render_markdown

    monkeypatch.setattr(perf_report, "SLOW_STATS_LIMIT", 2)
    for i in range(3):
        _slow(db_session, NOW - timedelta(minutes=10 + i), db_ms=2500)
    db_session.commit()

    data = overview(db_session, "1h", now=NOW, slow_ms=2000, flush_seconds=60)

    assert data["totals"]["slow"] == 3
    assert data["verdicts_basis"] == 2
    assert sum(data["verdicts"].values()) == 2
    assert "(по последним 2)" in render_markdown(data)
