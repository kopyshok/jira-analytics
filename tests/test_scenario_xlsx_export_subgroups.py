"""Выгрузка сценария команды с группами: у идей колонка «Группа» по тому же
правилу, что секции на экране (своя группа задачи, иначе группа исполнителя)."""

from io import BytesIO

from openpyxl import load_workbook

from app.models import (
    BacklogItem,
    Employee,
    EmployeeTeam,
    Issue,
    PlanningScenario,
    Project,
    Role,
    ScenarioAllocation,
    Team,
    TeamSubgroup,
)
from app.services.scenario_xlsx_export import ScenarioXlsxExporter
from tests.subgroup_fixtures import share

TEAM = "Бухгалтерия"


def _rows(ws) -> dict[str, dict[str, object]]:
    headers = [c.value for c in ws[2]]
    out = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] and str(row[0]).startswith("Σ"):
            continue
        out[row[1]] = dict(zip(headers, row))
    return out


def _seed(db, with_groups: bool = True):
    db.add(Role(code="dev", label="Разработчик", is_active=True, counts_in_planning=True))
    db.add(Team(id="t-1", name=TEAM, has_subgroups=with_groups))
    db.flush()
    db.add_all([
        TeamSubgroup(id="sg-a", team_id="t-1", name="Расчёты", sort_order=1),
        TeamSubgroup(id="sg-b", team_id="t-1", name="Интеграции", sort_order=2),
    ])
    emp = Employee(id="e-1", jira_account_id="acc-1", display_name="Иванов", role="dev",
                   is_active=True)
    db.add(emp)
    db.flush()
    db.add(EmployeeTeam(employee_id="e-1", team=TEAM, is_primary=True))
    db.add(share("e-1", TEAM, "sg-a"))
    proj = Project(id="p-1", jira_project_id="p1", key="P", name="P", is_active=True)
    db.add(proj)
    db.flush()
    issue_b = Issue(id="i-1", jira_issue_id="j1", key="P-1", summary="s", issue_type="Epic",
                    status="Open", project_id="p-1", team=TEAM, effective_subgroup_id="sg-b")
    db.add(issue_b)
    db.flush()
    own_group = BacklogItem(title="Своя группа задачи", priority=1, estimate_dev_hours=10,
                            issue_id="i-1", assignee_employee_id="e-1")
    by_assignee = BacklogItem(title="Группа исполнителя", priority=2, estimate_dev_hours=20,
                              assignee_employee_id="e-1")
    no_group = BacklogItem(title="Без исполнителя", priority=3, estimate_dev_hours=30)
    excluded = BacklogItem(title="Не вошла", priority=4, estimate_dev_hours=40,
                           assignee_employee_id="e-1")
    db.add_all([own_group, by_assignee, no_group, excluded])
    db.flush()
    sc = PlanningScenario(id="sc-1", name="Q4", year=2026, quarter="Q4", team=TEAM,
                          status="draft")
    db.add(sc)
    db.flush()
    for item, inc in ((own_group, True), (by_assignee, True), (no_group, True), (excluded, False)):
        db.add(ScenarioAllocation(scenario_id="sc-1", backlog_item_id=item.id,
                                  included_flag=inc, planned_hours=item.estimate_dev_hours))
    db.commit()


def test_group_column_follows_screen_rule(db_session):
    _seed(db_session)
    wb = load_workbook(BytesIO(ScenarioXlsxExporter(db_session, "sc-1").build()))

    included = _rows(wb["Включено"])
    assert included["Своя группа задачи"]["Группа"] == "Интеграции"
    assert included["Группа исполнителя"]["Группа"] == "Расчёты"
    assert included["Без исполнителя"]["Группа"] == "Без группы"
    assert _rows(wb["Не вошло"])["Не вошла"]["Группа"] == "Расчёты"


def test_totals_still_numeric_with_group_column(db_session):
    _seed(db_session)
    ws = load_workbook(BytesIO(ScenarioXlsxExporter(db_session, "sc-1").build()))["Включено"]
    headers = [c.value for c in ws[2]]
    total = next(r for r in ws.iter_rows(min_row=3, values_only=True)
                 if r[0] and str(r[0]).startswith("Σ"))
    assert total[headers.index("План, ч")] == 60
    assert total[headers.index("Группа")] in ("", None)


def test_team_without_groups_has_no_group_column(db_session):
    _seed(db_session, with_groups=False)
    wb = load_workbook(BytesIO(ScenarioXlsxExporter(db_session, "sc-1").build()))
    assert "Группа" not in [c.value for c in wb["Включено"][2]]
    assert "Группа" not in [c.value for c in wb["Не вошло"][2]]
