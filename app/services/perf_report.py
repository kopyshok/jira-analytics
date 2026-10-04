"""Чтение замеров быстродействия за период — для экрана админа и выгрузок.

Агрегаты суммируются в базе (`GROUP BY`): каждый сброс пишет свои строки, и
несколько строк одной минуты и пути — нормальное состояние.

Вывод по медленному запросу (порядок проверки):
1. сервер загружен (≥ 85%), а наш процесс — малая доля (< 30% машины) —
   ресурсы занял кто-то другой;
2. больше половины времени запроса — обращения к базе — долгая работа с базой;
3. процессор нашего процесса был занят за время запроса (≥ 70% ядра) или
   в ту минуту (≥ 80% ядра) — медленный наш код;
4. иначе — ожидание: не процессор и не база (например, ответ Jira).

ponytail: снимки нескольких процессов (если сервис поедет в несколько
воркеров) усредняются вместе, а не складываются по процессу — при одном
воркере это одно и то же.
"""
from __future__ import annotations

import bisect
from collections import Counter
from datetime import datetime, timedelta, timezone
from typing import Any, Optional

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.perf import HIST_SIZE, p95_from_hist
from app.models.perf import PERF_HIST_COLUMNS, PerfMinute, PerfServerSnapshot, PerfSlowRequest
from app.models.user import User

#: Период → (длина, размер точки графика в минутах).
PERIODS: dict[str, tuple[timedelta, int]] = {
    "1h": (timedelta(hours=1), 1),
    "24h": (timedelta(hours=24), 15),
    "7d": (timedelta(days=7), 60),
    "30d": (timedelta(days=30), 360),
}
PERIOD_LABELS = {
    "1h": "последний час",
    "24h": "последние сутки",
    "7d": "последние 7 дней",
    "30d": "последние 30 дней",
}

HOST_BUSY = 85.0
PROCESS_SMALL_SHARE = 30.0
DB_SHARE_HIGH = 0.5
REQUEST_CPU_HIGH = 70.0
PROCESS_CORE_HIGH = 80.0

VERDICT_LABELS = {
    "other_load": "Похоже, ресурсы сервера занял кто-то другой",
    "database": "Долгая работа с базой",
    "our_code": "Медленный наш код",
    "waiting": "Ожидание: не процессор и не база (например, ответ Jira)",
}

BOTTLENECKS_LIMIT = 50
SLOW_LIST_LIMIT = 200
SLOW_STATS_LIMIT = 10000

#: Первый сегмент пути после /api/v1/ → раздел, как он называется в меню.
SECTION_LABELS = {
    "admin": "Администрирование",
    "ai-status": "AI",
    "analytics": "Аналитика",
    "auth": "Вход",
    "backlog": "Целевые задачи",
    "capacity": "Ресурсы",
    "categories": "Категории работ",
    "desk": "Стол аналитика",
    "employees": "Сотрудники",
    "events": "Обновления на экране",
    "executive": "Сводка для руководителя",
    "exports": "Выгрузки",
    "feedback": "Обратная связь",
    "hierarchy-rules": "Правила иерархии",
    "issues": "Категории задач",
    "jira": "Синхронизация",
    "kpi": "KPI",
    "kpi-settings": "Настройки KPI",
    "llm": "AI",
    "mandatory-work-types": "Виды работ",
    "mapping": "Категории работ",
    "onboarding": "Первые шаги",
    "planning": "Сценарии",
    "production-calendar": "Производственный календарь",
    "projects": "Проекты",
    "release-notes": "Что нового",
    "resource-planning": "Ресурсное планирование",
    "roles": "Роли",
    "scope": "Проекты в scope",
    "settings": "Настройки",
    "sync": "Синхронизация",
    "team-desk": "Стол тимлида",
    "teams": "Команды",
    "themes": "Оформление",
    "ui-config": "Оформление",
    "usage": "Использование",
    "users": "Профиль",
    "work-desks": "Стол аналитика",
    "work-type-report": "Тематический отчёт",
}


def section_label(route: str) -> str:
    """Раздел по шаблону пути; незнакомый — первый сегмент как есть."""
    rest = route.split("/api/v1/", 1)[-1].lstrip("/")
    head = rest.split("/", 1)[0]
    return SECTION_LABELS.get(head, head or route)


def verdict(
    *,
    host_cpu: Optional[float],
    process_share: Optional[float],
    process_core: Optional[float],
    duration_ms: float,
    db_ms: float,
    cpu_ms: float,
) -> str:
    """Вероятная причина медленного запроса (см. порядок в описании модуля)."""
    if (
        host_cpu is not None and process_share is not None
        and host_cpu >= HOST_BUSY and process_share < PROCESS_SMALL_SHARE
    ):
        return "other_load"
    if duration_ms > 0 and db_ms / duration_ms >= DB_SHARE_HIGH:
        return "database"
    if duration_ms > 0 and cpu_ms / duration_ms * 100 >= REQUEST_CPU_HIGH:
        return "our_code"
    if process_core is not None and process_core >= PROCESS_CORE_HIGH:
        return "our_code"
    return "waiting"


def iso(dt: datetime) -> str:
    """Наивное UTC из базы → ISO с поясом, чтобы браузер не принял его за местное."""
    return dt.replace(tzinfo=timezone.utc).isoformat()


def _floor(dt: datetime, minutes: int) -> datetime:
    epoch = datetime(1970, 1, 1)
    step = minutes * 60
    seconds = int((dt - epoch).total_seconds()) // step * step
    return epoch + timedelta(seconds=seconds)


def _hist_sums() -> list[Any]:
    return [func.sum(getattr(PerfMinute, c)) for c in PERF_HIST_COLUMNS]


def _bottlenecks(db: Session, start: datetime) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = (
        db.query(
            PerfMinute.method, PerfMinute.route,
            func.sum(PerfMinute.count), func.sum(PerfMinute.total_ms), func.max(PerfMinute.max_ms),
            func.sum(PerfMinute.errors_5xx), func.sum(PerfMinute.db_count),
            func.sum(PerfMinute.db_ms), *_hist_sums(),
        )
        .filter(PerfMinute.minute >= start)
        .group_by(PerfMinute.method, PerfMinute.route)
        .all()
    )
    items = []
    total_hist = [0] * HIST_SIZE
    totals = {"requests": 0, "errors_5xx": 0, "total_ms": 0.0, "max_ms": 0.0}
    for method, route, count, total_ms, max_ms, errors, db_count, db_ms, *hist in rows:
        count, total_ms, max_ms = int(count or 0), float(total_ms or 0), float(max_ms or 0)
        if count <= 0:
            continue
        hist = [int(h or 0) for h in hist]
        total_hist = [a + b for a, b in zip(total_hist, hist)]
        totals["requests"] += count
        totals["errors_5xx"] += int(errors or 0)
        totals["total_ms"] += total_ms
        totals["max_ms"] = max(totals["max_ms"], max_ms)
        items.append({
            "method": method,
            "route": route,
            "section": section_label(route),
            "calls": count,
            "total_ms": total_ms,
            "avg_ms": total_ms / count,
            "p95_ms": p95_from_hist(hist, max_ms),
            "max_ms": max_ms,
            "errors_5xx": int(errors or 0),
            "db_avg_count": int(db_count or 0) / count,
            "db_share": float(db_ms or 0) / total_ms if total_ms > 0 else 0.0,
        })
    items.sort(key=lambda r: -r["total_ms"])
    totals["p95_ms"] = p95_from_hist(total_hist, totals["max_ms"])
    totals["avg_ms"] = totals["total_ms"] / totals["requests"] if totals["requests"] else 0.0
    return items[:BOTTLENECKS_LIMIT], totals


def _snapshots(db: Session, start: datetime) -> list[PerfServerSnapshot]:
    return (
        db.query(PerfServerSnapshot)
        .filter(PerfServerSnapshot.at >= start)
        .order_by(PerfServerSnapshot.at)
        .all()
    )


def _share(s: PerfServerSnapshot) -> float:
    """Доля процесса от всей машины, % — в тех же единицах, что загрузка сервера."""
    return s.process_cpu_percent / max(s.cpu_count, 1)


class _SnapshotIndex:
    """Поиск снимка, интервал которого покрывает момент запроса."""

    def __init__(self, snaps: list[PerfServerSnapshot], tolerance: timedelta) -> None:
        self.snaps = snaps
        self.ats = [s.at for s in snaps]
        self.tolerance = tolerance

    def find(self, at: datetime) -> Optional[PerfServerSnapshot]:
        i = bisect.bisect_left(self.ats, at)
        # Снимок пишется в конце интервала — сначала ближайший после запроса.
        if i < len(self.snaps) and self.ats[i] - at <= self.tolerance:
            return self.snaps[i]
        if i > 0 and at - self.ats[i - 1] <= self.tolerance:
            return self.snaps[i - 1]
        return None


def _load(snap: Optional[PerfServerSnapshot]) -> dict[str, Any]:
    if snap is None:
        return {
            "host_cpu": None, "process_cpu": None, "process_cpu_core": None,
            "host_memory": None, "process_memory_mb": None, "threads": None,
            "db_pool_in_use": None, "db_pool_size": None, "requests_in_flight": None,
        }
    return {
        "host_cpu": snap.host_cpu_percent,
        "process_cpu": _share(snap),
        "process_cpu_core": snap.process_cpu_percent,
        "host_memory": snap.host_memory_percent,
        "process_memory_mb": snap.process_memory_mb,
        "threads": snap.threads,
        "db_pool_in_use": snap.db_pool_in_use,
        "db_pool_size": snap.db_pool_size,
        "requests_in_flight": snap.requests_in_flight,
    }


def _verdict_for(slow: Any, snap: Optional[PerfServerSnapshot]) -> str:
    return verdict(
        host_cpu=snap.host_cpu_percent if snap else None,
        process_share=_share(snap) if snap else None,
        process_core=snap.process_cpu_percent if snap else None,
        duration_ms=slow.duration_ms,
        db_ms=slow.db_ms,
        cpu_ms=slow.cpu_ms,
    )


def overview(
    db: Session,
    period: str,
    *,
    now: Optional[datetime] = None,
    slow_ms: float,
    flush_seconds: float,
    slow_limit: int = SLOW_LIST_LIMIT,
) -> dict[str, Any]:
    """Всё для экрана за период: итог, ряд для графика, узкие места, медленные."""
    if period not in PERIODS:
        raise ValueError(f"Неизвестный период: {period}")
    length, bucket_minutes = PERIODS[period]
    now = now or datetime.utcnow()
    start = _floor(now - length, 1)
    bucket = timedelta(minutes=bucket_minutes)

    bottlenecks, totals = _bottlenecks(db, start)
    snaps = _snapshots(db, start)
    index = _SnapshotIndex(snaps, tolerance=timedelta(seconds=max(2 * flush_seconds, 120)))

    # --- ряд для графика: пустые точки тоже, чтобы ось времени была сплошной
    points: dict[datetime, dict[str, Any]] = {}
    t = _floor(start, bucket_minutes)
    while t <= now:
        points[t] = {"hist": [0] * HIST_SIZE, "count": 0, "max_ms": 0.0,
                     "host": [], "proc": [], "slow": Counter()}
        t += bucket

    def point(at: datetime) -> Optional[dict[str, Any]]:
        return points.get(_floor(at, bucket_minutes))

    per_minute = (
        db.query(PerfMinute.minute, func.sum(PerfMinute.count), func.max(PerfMinute.max_ms),
                 *_hist_sums())
        .filter(PerfMinute.minute >= start)
        .group_by(PerfMinute.minute)
        .all()
    )
    for minute, count, max_ms, *hist in per_minute:
        p = point(minute)
        if p is None:
            continue
        p["count"] += int(count or 0)
        p["max_ms"] = max(p["max_ms"], float(max_ms or 0))
        p["hist"] = [a + int(b or 0) for a, b in zip(p["hist"], hist)]
    for s in snaps:
        p = point(s.at)
        if p is not None:
            p["host"].append(s.host_cpu_percent)
            p["proc"].append(_share(s))

    # --- медленные: все за период для счёта причин, последние — с деталями
    slow_cols = (
        db.query(PerfSlowRequest.at, PerfSlowRequest.duration_ms, PerfSlowRequest.db_ms,
                 PerfSlowRequest.cpu_ms)
        .filter(PerfSlowRequest.at >= start)
        .order_by(PerfSlowRequest.at.desc())
        .limit(SLOW_STATS_LIMIT)
        .all()
    )
    verdicts: Counter = Counter()
    for row in slow_cols:
        v = _verdict_for(row, index.find(row.at))
        verdicts[v] += 1
        p = point(row.at)
        if p is not None:
            p["slow"][v] += 1

    series = []
    for t, p in points.items():
        top = p["slow"].most_common(1)
        series.append({
            "t": iso(t),
            "requests": p["count"],
            "p95_ms": p95_from_hist(p["hist"], p["max_ms"]) if p["count"] else None,
            "max_ms": p["max_ms"] if p["count"] else None,
            "slow": sum(p["slow"].values()),
            "host_cpu": sum(p["host"]) / len(p["host"]) if p["host"] else None,
            "host_cpu_max": max(p["host"]) if p["host"] else None,
            "process_cpu": sum(p["proc"]) / len(p["proc"]) if p["proc"] else None,
            "verdict": top[0][0] if top else None,
            "verdict_label": VERDICT_LABELS[top[0][0]] if top else None,
        })

    slow_rows = (
        db.query(PerfSlowRequest)
        .filter(PerfSlowRequest.at >= start)
        .order_by(PerfSlowRequest.at.desc())
        .limit(slow_limit)
        .all()
    )
    user_ids = {r.user_id for r in slow_rows if r.user_id}
    names = {
        u.id: (u.display_name or u.email)
        for u in db.query(User).filter(User.id.in_(user_ids)).all()
    } if user_ids else {}
    slow = []
    for r in slow_rows:
        snap = index.find(r.at)
        v = _verdict_for(r, snap)
        slow.append({
            "id": r.id,
            "at": iso(r.at),
            "method": r.method,
            "route": r.route,
            "section": section_label(r.route),
            "path": r.path,
            "query": r.query,
            "status_code": r.status_code,
            "duration_ms": r.duration_ms,
            "db_count": r.db_count,
            "db_ms": r.db_ms,
            "cpu_ms": r.cpu_ms,
            "user": names.get(r.user_id or ""),
            "top_queries": r.top_queries or [],
            **_load(snap),
            "verdict": v,
            "verdict_label": VERDICT_LABELS[v],
        })

    return {
        "period": period,
        "period_label": PERIOD_LABELS[period],
        "start": iso(start),
        "end": iso(now),
        "bucket_minutes": bucket_minutes,
        "slow_ms": slow_ms,
        "totals": {**totals, "slow": len(slow_cols)},
        "verdicts": dict(verdicts),
        "verdict_labels": VERDICT_LABELS,
        "series": series,
        "bottlenecks": bottlenecks,
        "slow": slow,
        "load": _load_summary(snaps),
    }


def _load_summary(snaps: list[PerfServerSnapshot]) -> Optional[dict[str, Any]]:
    """Нагрузка сервера за период: средние, пики и «тяжёлые» снимки."""
    if not snaps:
        return None
    host = [s.host_cpu_percent for s in snaps]
    share = [_share(s) for s in snaps]
    pool_sizes = [s.db_pool_size for s in snaps if s.db_pool_size is not None]
    pool_used = [s.db_pool_in_use for s in snaps if s.db_pool_in_use is not None]
    return {
        "snapshots": len(snaps),
        "host_cpu_avg": sum(host) / len(host),
        "host_cpu_max": max(host),
        "process_cpu_avg": sum(share) / len(share),
        "process_cpu_max": max(share),
        "host_memory_max": max(s.host_memory_percent for s in snaps),
        "process_memory_mb_max": max(s.process_memory_mb for s in snaps),
        "threads_max": max(s.threads for s in snaps),
        "db_pool_in_use_max": max(pool_used) if pool_used else None,
        "db_pool_size": max(pool_sizes) if pool_sizes else None,
        "requests_in_flight_max": max(s.requests_in_flight for s in snaps),
        "busy_snapshots": sum(1 for h in host if h >= HOST_BUSY),
        "busy_by_others_snapshots": sum(
            1 for h, p in zip(host, share) if h >= HOST_BUSY and p < PROCESS_SMALL_SHARE
        ),
    }


def snapshots_for(
    db: Session, period: str, *, now: Optional[datetime] = None,
) -> list[dict[str, Any]]:
    """Все снимки сервера за период — для листа Excel."""
    length, _ = PERIODS[period]
    start = _floor((now or datetime.utcnow()) - length, 1)
    return [{"at": iso(s.at), **_load(s)} for s in _snapshots(db, start)]


# --- Выгрузки ------------------------------------------------------------------

REPORT_SLOW_LIMIT = 50
REPORT_BOTTLENECKS_LIMIT = 20


def _tz_label(tz_offset_min: int) -> str:
    sign = "+" if tz_offset_min >= 0 else "-"
    h, m = divmod(abs(tz_offset_min), 60)
    return f"UTC{sign}{h:02d}:{m:02d}"


def _local(value: str, tz_offset_min: int) -> datetime:
    """ISO из ответа → наивное местное время (Excel не принимает пояс)."""
    tz = timezone(timedelta(minutes=tz_offset_min))
    return datetime.fromisoformat(value).astimezone(tz).replace(tzinfo=None)


def _dt(value: str, tz_offset_min: int, seconds: bool = True) -> str:
    fmt = "%d.%m.%Y %H:%M:%S" if seconds else "%d.%m.%Y %H:%M"
    return _local(value, tz_offset_min).strftime(fmt)


def _n(x: Optional[float]) -> str:
    return "—" if x is None else f"{round(x):,}".replace(",", " ")


def _pct(x: Optional[float]) -> str:
    return "—" if x is None else f"{round(x)}%"


def _load_line(s: dict[str, Any]) -> str:
    if s.get("host_cpu") is None:
        return "нет снимка нагрузки рядом с этим моментом"
    pool = (
        f"{s['db_pool_in_use']}/{s['db_pool_size']}"
        if s.get("db_pool_size") is not None else "нет данных"
    )
    return (
        f"процессор сервера {_pct(s['host_cpu'])}, наш процесс {_pct(s['process_cpu_core'])} ядра "
        f"({_pct(s['process_cpu'])} машины), память сервера {_pct(s['host_memory'])}, "
        f"память процесса {_n(s['process_memory_mb'])} МБ, потоков {s['threads']}, "
        f"подключения к базе {pool}, запросов в работе до {s['requests_in_flight']}"
    )


def render_markdown(data: dict[str, Any], *, tz_offset_min: int = 0) -> str:
    """Отчёт для разработки: итог, узкие места, медленные с деталями, нагрузка сервера."""
    tz = tz_offset_min
    totals, load = data["totals"], data["load"]
    out: list[str] = [
        "# Быстродействие сервиса — отчёт для разработки",
        "",
        f"Период: {data['period_label']} — с {_dt(data['start'], tz, False)} "
        f"по {_dt(data['end'], tz, False)} ({_tz_label(tz)}).",
        f"Порог медленного запроса: {_n(data['slow_ms'])} мс.",
        "",
        "## Итог",
        "",
        f"- Запросов к API: {_n(totals['requests'])}; медленных: {_n(totals['slow'])}; "
        f"ошибок сервера (5xx): {_n(totals['errors_5xx'])}.",
        f"- Среднее время ответа: {_n(totals['avg_ms'])} мс; 95-й процентиль: "
        f"{_n(totals['p95_ms'])} мс; максимум: {_n(totals['max_ms'])} мс.",
    ]
    if data["verdicts"]:
        out.append("- Вероятные причины медленных запросов:")
        for code, n in sorted(data["verdicts"].items(), key=lambda kv: -kv[1]):
            out.append(f"  - {VERDICT_LABELS[code]} — {n}")

    out += ["", f"## Узкие места (по суммарному времени, топ {REPORT_BOTTLENECKS_LIMIT})", ""]
    if data["bottlenecks"]:
        out += [
            "| # | Раздел | Запрос | Вызовов | Всего, с | Среднее, мс | p95, мс | Макс, мс "
            "| Обращений к базе (сред.) | Доля базы | 5xx |",
            "|---|---|---|---|---|---|---|---|---|---|---|",
        ]
        for i, b in enumerate(data["bottlenecks"][:REPORT_BOTTLENECKS_LIMIT], 1):
            out.append(
                f"| {i} | {b['section']} | `{b['method']} {b['route']}` | {_n(b['calls'])} "
                f"| {b['total_ms'] / 1000:.1f} | {_n(b['avg_ms'])} | {_n(b['p95_ms'])} "
                f"| {_n(b['max_ms'])} | {b['db_avg_count']:.1f} | {_pct(b['db_share'] * 100)} "
                f"| {b['errors_5xx']} |"
            )
    else:
        out.append("Запросов за период нет.")

    slowest = sorted(data["slow"], key=lambda s: -s["duration_ms"])[:REPORT_SLOW_LIMIT]
    out += ["", f"## Медленные запросы (самые долгие {len(slowest)} из {totals['slow']})", ""]
    if not slowest:
        out.append("Медленных запросов за период нет.")
    for i, s in enumerate(slowest, 1):
        address = s["path"] + (f"?{s['query']}" if s["query"] else "")
        db_share = s["db_ms"] / s["duration_ms"] * 100 if s["duration_ms"] else 0
        cpu = s["cpu_ms"] / s["duration_ms"] * 100 if s["duration_ms"] else 0
        out += [
            f"### {i}. {_dt(s['at'], tz)} — `{s['method']} {s['route']}` — "
            f"{_n(s['duration_ms'])} мс",
            "",
            f"- Раздел: {s['section']}; пользователь: {s['user'] or '—'}; ответ: {s['status_code']}",
            f"- Адрес: `{address}`",
            f"- База: обращений {s['db_count']}, {_n(s['db_ms'])} мс ({_pct(db_share)} времени)",
            f"- Процессор нашего процесса за время запроса: {_pct(cpu)} ядра",
            f"- Нагрузка сервера в ту минуту: {_load_line(s)}",
            f"- **Вывод: {s['verdict_label']}**",
        ]
        if s["top_queries"]:
            out += ["", "Самые долгие обращения к базе:", ""]
            for j, q in enumerate(s["top_queries"], 1):
                out += [f"{j}. {_n(q['ms'])} мс", "   ```sql", f"   {q['sql']}", "   ```"]
        out.append("")

    if out[-1]:
        out.append("")
    out += ["## Нагрузка сервера", ""]
    if load is None:
        out.append("Снимков нагрузки за период нет.")
    else:
        pool = (
            f"{load['db_pool_in_use_max']} из {load['db_pool_size']}"
            if load["db_pool_size"] is not None else "нет данных"
        )
        out += [
            f"- Процессор сервера: в среднем {_pct(load['host_cpu_avg'])}, "
            f"пик {_pct(load['host_cpu_max'])}.",
            f"- Наш процесс (доля машины): в среднем {_pct(load['process_cpu_avg'])}, "
            f"пик {_pct(load['process_cpu_max'])}.",
            f"- Память: сервер до {_pct(load['host_memory_max'])}, процесс до "
            f"{_n(load['process_memory_mb_max'])} МБ; потоков до {load['threads_max']}.",
            f"- Подключения к базе: занято до {pool}; запросов в работе одновременно до "
            f"{load['requests_in_flight_max']}.",
            f"- Снимков с загрузкой сервера от {round(HOST_BUSY)}%: {load['busy_snapshots']} "
            f"из {load['snapshots']}; из них наш процесс меньше {round(PROCESS_SMALL_SHARE)}% "
            f"машины (сервер занят не нами): {load['busy_by_others_snapshots']}.",
        ]
    return "\n".join(out) + "\n"


def render_xlsx(
    data: dict[str, Any], snapshots: list[dict[str, Any]], *, tz_offset_min: int = 0,
) -> bytes:
    """Excel: узкие места, медленные запросы, снимки сервера."""
    from io import BytesIO

    from openpyxl import Workbook  # type: ignore[import-untyped]
    from openpyxl.styles import Font  # type: ignore[import-untyped]

    tz = tz_offset_min
    wb = Workbook()

    def fill(ws: Any, header: list[str], rows: list[list[Any]], widths: list[int]) -> None:
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    def r0(x: Optional[float]) -> Optional[float]:
        return None if x is None else round(x, 1)

    ws = wb.active
    ws.title = "Узкие места"
    fill(ws, [
        "Раздел", "Метод", "Путь", "Вызовов", "Всего, с", "Среднее, мс", "95-й процентиль, мс",
        "Максимум, мс", "Обращений к базе в среднем", "Доля времени в базе, %", "Ошибок 5xx",
    ], [
        [b["section"], b["method"], b["route"], b["calls"], round(b["total_ms"] / 1000, 1),
         round(b["avg_ms"]), round(b["p95_ms"]), round(b["max_ms"]), round(b["db_avg_count"], 1),
         round(b["db_share"] * 100), b["errors_5xx"]]
        for b in data["bottlenecks"]
    ], [24, 8, 50, 10, 10, 12, 14, 12, 14, 14, 10])

    time_header = f"Время ({_tz_label(tz)})"
    fill(wb.create_sheet("Медленные запросы"), [
        time_header, "Раздел", "Метод", "Путь", "Адрес", "Параметры", "Пользователь", "Ответ",
        "Длительность, мс", "Обращений к базе", "Время в базе, мс",
        "Процессор процесса за запрос, мс", "Процессор сервера, %", "Наш процесс, % машины",
        "Вывод", "Самые долгие обращения к базе",
    ], [
        [_local(s["at"], tz), s["section"], s["method"], s["route"], s["path"], s["query"],
         s["user"] or "", s["status_code"], round(s["duration_ms"]), s["db_count"],
         round(s["db_ms"]), round(s["cpu_ms"]), r0(s["host_cpu"]), r0(s["process_cpu"]),
         s["verdict_label"],
         "\n\n".join(f"{round(q['ms'])} мс: {q['sql']}" for q in s["top_queries"])]
        for s in data["slow"]
    ], [20, 22, 8, 44, 40, 30, 22, 8, 14, 12, 14, 16, 14, 14, 40, 80])

    fill(wb.create_sheet("Снимки сервера"), [
        time_header, "Процессор сервера, %", "Наш процесс, % машины", "Наш процесс, % ядра",
        "Память сервера, %", "Память процесса, МБ", "Потоков", "Подключений к базе занято",
        "Подключений к базе всего", "Запросов в работе (пик)",
    ], [
        [_local(s["at"], tz), r0(s["host_cpu"]), r0(s["process_cpu"]), r0(s["process_cpu_core"]),
         r0(s["host_memory"]), round(s["process_memory_mb"]), s["threads"], s["db_pool_in_use"],
         s["db_pool_size"], s["requests_in_flight"]]
        for s in snapshots
    ], [20, 14, 14, 14, 14, 14, 10, 14, 14, 14])

    for sheet in wb.worksheets[1:]:
        for (cell,) in sheet.iter_rows(min_row=2, max_col=1):
            cell.number_format = "DD.MM.YYYY HH:MM:SS"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
